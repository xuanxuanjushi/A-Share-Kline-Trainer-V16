"""Excel reports for the desktop app; no dependency on the authoring environment."""
from __future__ import annotations

import math
import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.data_source import NumData, NumVal
from openpyxl.chart.text import RichText, Text
from openpyxl.drawing.text import CharacterProperties, Paragraph, RegularTextRun
from openpyxl.descriptors import Typed
from openpyxl.descriptors.nested import NestedBool
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.table import Table, TableStyleInfo


_LABEL_TEXT_OFF = dict(showLegendKey=False, showVal=False, showCatName=False,
                       showSerName=False, showPercent=False, showBubbleSize=False,
                       showLeaderLines=False)


class _TextDataLabel(DataLabel):
    """OOXML supports rich label text; openpyxl's base label omits this field."""
    tx = Typed(expected_type=Text, allow_none=True)
    delete = NestedBool(allow_none=True)
    __elements__ = ('idx', 'delete', 'tx', *DataLabel.__elements__[1:])

    def __init__(self, idx, text=None):
        super().__init__(idx=idx, dLblPos='b', **_LABEL_TEXT_OFF)
        self.delete = text is None
        self.tx = (Text(rich=RichText(p=[Paragraph(r=[RegularTextRun(
            rPr=CharacterProperties(solidFill='D64045', sz=900), t=text)])])) if text is not None else None)


_TextDataLabel.__nested__ = (*DataLabel.__nested__, 'delete')


def _value(value):
    if isinstance(value, float) and not math.isfinite(value):
        return '∞' if value > 0 else '-∞' if value < 0 else None
    return value


def _cell(cell, value):
    cell.value = _value(value)
    if isinstance(cell.value, str):
        cell.data_type = 's'  # Stock names and external text must never become formulas.


def _table(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    for row_index, row in enumerate([headers, *rows], 1):
        for column, value in enumerate(row, 1):
            cell = ws.cell(row_index, column)
            _cell(cell, value)
            cell.font = Font(name='微软雅黑', size=10, color='243746')
            cell.alignment = Alignment(vertical='center')
            if isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd'
            elif isinstance(value, float):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
        ws.row_dimensions[row_index].height = 22
    for i, header in enumerate(headers, 1):
        cell = ws.cell(1, i)
        cell.fill = PatternFill('solid', fgColor='16324F')
        cell.font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
        ws.column_dimensions[get_column_letter(i)].width = min(30, max(15, len(header) * 2 + 2))
        if str(header).endswith('%'):
            for cells in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cells[0].number_format = '0.00"%";[Red]-0.00"%"'
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    if rows:
        table = Table(displayName=f'Data{len(wb.worksheets)}', ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    ws.auto_filter.ref = ws.dimensions
    return ws


def _daily_rows(record):
    by_date = {}
    for point in record.equity_points:
        if not math.isfinite(point.total_asset) or point.total_asset <= 0:
            continue
        # Prefer the last close, or the latest available open for an unfinished day.
        previous = by_date.get(point.date)
        if previous is None or point.node.value == 'close' or previous.node.value != 'close':
            by_date[point.date] = point
    initial = record.start_cash
    peak = initial if initial > 0 else 0
    result = []
    for date, point in sorted(by_date.items()):
        peak = max(peak, point.total_asset)
        result.append([record.index, datetime.strptime(date, '%Y-%m-%d'), point.total_asset,
                       initial if initial > 0 else None,
                       point.total_asset / initial - 1 if initial > 0 else None,
                       point.total_asset / peak - 1, peak - point.total_asset,
                       '收盘' if point.node.value == 'close' else '开盘（未收盘）'])
    return result


def _chart(title, ws, start, end, column, color):
    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.width, chart.height = 25, 12
    chart.x_axis = DateAxis(crossAx=100)
    chart.y_axis.crossAx = chart.x_axis.axId
    chart.x_axis.majorTimeUnit = 'months'
    chart.x_axis.majorUnit = max(1, (end - start) // 180)
    if end > start:
        chart.x_axis.scaling.min = to_excel(ws.cell(start, 2).value)
        chart.x_axis.scaling.max = to_excel(ws.cell(end, 2).value)
    chart.x_axis.number_format = 'yyyy-mm'
    chart.x_axis.title = '交易日期'
    chart.add_data(Reference(ws, min_col=column, min_row=start, max_row=end))
    chart.set_categories(Reference(ws, min_col=2, min_row=start, max_row=end))
    chart.series[0].graphicalProperties.line.solidFill = color
    chart.series[0].graphicalProperties.line.width = 22000
    chart.series[0].val.numRef.numCache = NumData(
        ptCount=end - start + 1,
        pt=[NumVal(idx=i, v=ws.cell(row, column).value) for i, row in enumerate(range(start, end + 1))
            if ws.cell(row, column).value is not None])
    chart.series[0].cat.numRef.numCache = NumData(
        formatCode='yyyy-mm-dd', ptCount=end - start + 1,
        pt=[NumVal(idx=i, v=to_excel(ws.cell(row, 2).value)) for i, row in enumerate(range(start, end + 1))])
    chart.display_blanks = 'gap'
    return chart


def _charts(target, daily, start, end, rows):
    assets = _chart('资金与回撤', daily, start, end, 3, '2474B5')
    assets.y_axis.title = '账户总资产（元）'
    assets.y_axis.numFmt = '#,##0'
    drawdown = _chart('回撤', daily, start, end, 6, 'D64045')
    drawdown.y_axis.axId = 200
    drawdown.y_axis.axPos = 'r'
    drawdown.y_axis.title = '较历史最高资产回撤'
    drawdown.y_axis.numFmt = '0%'
    drawdown.y_axis.crosses = 'max'
    drawdown.y_axis.scaling.max = 0
    drawdown.y_axis.scaling.min = min(-0.01, min(row[5] for row in rows) * 1.35)
    drawdown.y_axis.majorGridlines = None
    # One label at the deepest point of each completed/ongoing drawdown episode.
    troughs, trough = [], None
    for index, row in enumerate(rows):
        if row[5] < -1e-10:
            if trough is None or row[5] < rows[trough][5]:
                trough = index
        elif trough is not None:
            troughs.append(trough)
            trough = None
    if trough is not None:
        troughs.append(trough)
    selected_labels = {}
    # Dense daily histories can have hundreds of tiny drawdowns. Label the 8
    # largest episodes, while retaining every amount and ratio in the data sheet.
    for index in sorted(sorted(troughs, key=lambda i: rows[i][5])[:8]):
        row = rows[index]
        selected_labels[index] = f'回撤 {row[6]:,.2f} 元（{abs(row[5]):.2%}）'
    # Excel/WPS may enable automatic labels when defaults are omitted. Disable
    # every default field and explicitly hide every unselected observation.
    labels = [_TextDataLabel(index, selected_labels.get(index)) for index in range(len(rows))]
    drawdown.series[0].dLbls = DataLabelList(dLbl=labels, **_LABEL_TEXT_OFF)
    from openpyxl.chart.series import SeriesLabel
    assets.series[0].tx = SeriesLabel(v='资金（元）')
    drawdown.series[0].tx = SeriesLabel(v='回撤比例（红色，右轴）')
    assets += drawdown
    target.add_chart(assets, 'D5')
    returns = _chart('累计收益率', daily, start, end, 5, '2474B5')
    returns.y_axis.numFmt = '0%'
    returns.y_axis.title = '相对该轮期初资金'
    returns.legend = None
    target.add_chart(returns, 'D25')


def write_performance_workbook(path: Path, rows: list[dict], fields: list[str], records) -> None:
    wb = Workbook()
    overview = wb.active
    overview.title = '统计概览'
    overview.sheet_view.showGridLines = False
    overview.column_dimensions['A'].width = 27
    overview.column_dimensions['B'].width = 25
    overview.column_dimensions['C'].width = 3
    for column in range(4, 20):
        overview.column_dimensions[get_column_letter(column)].width = 12
    for row in range(1, 60):
        overview.row_dimensions[row].height = 23
    overview.merge_cells('A1:R2')
    overview['A1'] = '交易统计与资金回撤'
    overview['A1'].font = Font(name='微软雅黑', size=22, bold=True, color='FFFFFF')
    overview['A1'].fill = PatternFill('solid', fgColor='16324F')
    summary = rows[0] if rows else {}
    for index, (key, value) in enumerate(summary.items(), 4):
        _cell(overview.cell(index, 1), key)
        _cell(overview.cell(index, 2), value)
        overview.cell(index, 2).alignment = Alignment(horizontal='left', vertical='center')
        overview.cell(index, 1).font = Font(name='微软雅黑', size=10, color='526577')
        overview.cell(index, 2).font = Font(name='微软雅黑', size=11, bold=True, color='16324F')
        if isinstance(value, (int, float)):
            overview.cell(index, 2).number_format = '0.00"%"' if key.endswith('%') else '#,##0.00'
            if key.endswith(('次数', '段数', '数量', '交易日', '样本日')):
                overview.cell(index, 2).number_format = '#,##0'

    for level, title in [('股票轮次', '股票轮次'), ('交易分段', '交易分段'), ('成交明细', '成交明细')]:
        subset = [row for row in rows if row.get('记录层级') == level]
        headers = [field for field in fields if field != '记录层级' and any(field in row for row in subset)]
        _table(wb, title, headers or ['说明'], [[row.get(f) for f in headers] for row in subset])
    _table(wb, '原始统计', fields, [[row.get(f) for f in fields] for row in rows])
    groups = [(record, _daily_rows(record)) for record in records]
    daily = _table(wb, '每日资产', ['轮次', '交易日期', '账户总资产', '期初资金', '累计收益率', '回撤比例', '回撤金额', '记录阶段'],
                   [row for _, data in groups for row in data])
    for row in daily.iter_rows(min_row=2):
        row[4].number_format = row[5].number_format = '0.00%'
        row[5].font = row[6].font = Font(name='微软雅黑', color='D64045')
    _table(wb, '资产原始记录', ['轮次', '记录顺序', '日期', '阶段', '账户总资产', '股票代码'],
           [[record.index, i, point.date, point.node.value, point.total_asset, point.code]
            for record in records for i, point in enumerate(record.equity_points, 1)])
    overview['D3'] = ('资金蓝色；回撤红色。标注最大8段回撤的金额和比例；全部数值见每日资产。'
                       if len(records) == 1 else '各轮曲线分别展示，避免不同训练日期重叠或倒序造成误连。')
    if any(not data for _, data in groups):
        overview['D3'] = '部分轮次缺少历史资产记录，保留统计与成交数据，不补造资金曲线。'
    overview['D4'] = '按日期排序，取最后收盘记录；未收盘日取开盘。缺失交易日不补值。'
    start = 2
    for number, (record, data) in enumerate(groups, 1):
        if not data:
            continue
        target = overview if len(records) == 1 else wb.create_sheet(f'曲线-{number}')
        if target is not overview:
            target.sheet_view.showGridLines = False
            target['D1'] = f'第{record.index}轮 {record.code} {record.name}'
            target['D3'] = '蓝色为资金，红色为回撤；标注最大8段回撤。完整数值见每日资产。'
        _charts(target, daily, start, start + len(data) - 1, data)
        start += len(data)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix='.excel-', suffix='.xlsx', dir=path.parent)
    os.close(descriptor)
    try:
        # A managed stream closes even when chart serialization fails.
        with open(temporary, 'wb') as handle:
            wb.save(handle)
        os.replace(temporary, path)
    finally:
        wb.close()
        Path(temporary).unlink(missing_ok=True)
