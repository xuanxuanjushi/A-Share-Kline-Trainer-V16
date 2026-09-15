import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from stock_simulator.models import EquityPoint, RoundRecord, TradeNode
from stock_simulator.excel_export import write_performance_workbook


class ExcelExportTests(unittest.TestCase):
    def test_only_selected_trough_has_a_label_and_all_default_text_is_disabled(self):
        record = RoundRecord(1, 'sz000061', '测试', -30, start_cash=100,
            equity_points=[EquityPoint(f'2020-01-0{i}', TradeNode.CLOSE, value)
                           for i, value in enumerate([100, 90, 80, 70], 1)])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'labels.xlsx'
            write_performance_workbook(path, [{'记录层级':'区间汇总'}], ['记录层级'], [record])
            with zipfile.ZipFile(path) as archive:
                root = ET.fromstring(archive.read('xl/charts/chart1.xml'))
            ns = {'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
            group = root.find('.//c:dLbls', ns)
            for flag in ('showLegendKey', 'showVal', 'showCatName', 'showSerName', 'showPercent', 'showBubbleSize', 'showLeaderLines'):
                self.assertIsNotNone(group.find(f'c:{flag}', ns), flag)
                self.assertEqual(group.find(f'c:{flag}', ns).get('val'), '0')
            labels = group.findall('c:dLbl', ns)
            self.assertEqual(len(labels), 4)
            for index, label in enumerate(labels):
                self.assertEqual(label.find('c:idx', ns).get('val'), str(index))
                self.assertEqual(label.find('c:delete', ns).get('val'), '0' if index == 3 else '1')
            self.assertIn('回撤 30.00 元（30.00%）', ET.tostring(labels[-1], encoding='unicode'))

    def test_preserves_all_statistics_and_charts_ordered_daily_assets(self):
        points = [EquityPoint('2020-01-03', TradeNode.CLOSE, 90),
                  EquityPoint('2020-01-02', TradeNode.OPEN, 99),
                  EquityPoint('2020-01-02', TradeNode.CLOSE, 110)]
        record = RoundRecord(1, 'sz000061', '=not a formula', 0, start_cash=100, equity_points=points)
        fields = ['记录层级', '股票代码', '股票名称', '净盈亏', '成交价']
        rows = [{'记录层级': '区间汇总', '净盈亏': -10},
                {'记录层级': '成交明细', '股票代码': '000061', '股票名称': '=not a formula', '成交价': 9.5}]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'report.xlsx'
            write_performance_workbook(path, rows, fields, [record])
            wb = load_workbook(path)
            self.assertEqual(list(wb['原始统计'].values)[1:], [tuple(row.get(f) for f in fields) for row in rows])
            self.assertEqual(wb['原始统计']['C3'].data_type, 's')
            daily = list(wb['每日资产'].values)
            self.assertEqual([r[1].strftime('%Y-%m-%d') for r in daily[1:]], ['2020-01-02', '2020-01-03'])
            self.assertEqual([r[2] for r in daily[1:]], [110, 90])
            self.assertAlmostEqual(daily[2][4], -.1)
            self.assertAlmostEqual(daily[2][5], 90 / 110 - 1)
            self.assertEqual(wb['资产原始记录'].max_row, 4)
            self.assertEqual(len(wb['统计概览']._charts), 2)
            with zipfile.ZipFile(path) as archive:
                chart = archive.read('xl/charts/chart1.xml').decode('utf-8')
                self.assertIn('回撤 20.00 元（18.18%）', chart)
                self.assertIn('D64045', chart)
                self.assertIn('axId val="200"', chart)
                self.assertIn('numCache', chart)
                self.assertIn('formatCode>yyyy-mm-dd', chart)
            self.assertEqual(wb['成交明细'].freeze_panes, 'A2')
            wb.close()

    def test_overlapping_rounds_have_separate_charts_and_missing_points_are_not_invented(self):
        records = [RoundRecord(i, 'sz000061', 'name', 0, start_cash=100,
                     equity_points=[EquityPoint('2020-01-02', TradeNode.CLOSE, 100 + i)]) for i in (1, 2)]
        records.append(RoundRecord(3, 'sz000061', 'old', 0))
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'report.xlsx'
            write_performance_workbook(path, [{'记录层级':'区间汇总'}], ['记录层级'], records)
            wb = load_workbook(path)
            self.assertEqual(wb['每日资产'].max_row, 3)
            self.assertIn('曲线-1', wb.sheetnames)
            self.assertIn('曲线-2', wb.sheetnames)
            self.assertNotIn('曲线-3', wb.sheetnames)
            self.assertEqual(len(wb['统计概览']._charts), 0)
            self.assertIn('缺少', wb['统计概览']['D3'].value)
            wb.close()
